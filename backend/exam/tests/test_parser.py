from datetime import datetime, time
from io import BytesIO
from openpyxl import Workbook
from openpyxl.utils.datetime import MAC_EPOCH, to_excel
from unittest import TestCase
import pandas as pd
from parser.exam import ExamParser


class ExamParserTest(TestCase):
    def get_parser(
        self, schedule: pd.DataFrame, reminder_minutes: int = 15
    ) -> ExamParser:
        parser = ExamParser.__new__(ExamParser)
        parser.reminder_minutes = reminder_minutes
        parser.students = pd.DataFrame(
            {"学号": [23000001, 23000002], "课程编号": ["FT03P207", "OTHER"]}
        )
        parser.schedule = (schedule,)
        return parser

    def test_full_date_time_text_becomes_event(self) -> None:
        schedule = pd.DataFrame(
            {
                "课程编号": ["FT03P207", "OTHER"],
                "课程名称": ["金融风险管理", "其他课程"],
                "考试日期": ["2026-09-17", "2026-09-18"],
                "考试时间": [
                    "2026-09-17 10:10~2026-09-17 11:50",
                    "08:20~10:00",
                ],
                "考试地点": ["教室 101", "教室 102"],
                "备注": [pd.NA, ""],
            }
        )

        events = list(self.get_parser(schedule).get_exam(23000001))

        self.assertEqual(len(events), 1)
        self.assertEqual(events[0].summary, "金融风险管理")
        self.assertEqual(events[0].dtstart, datetime(2026, 9, 17, 10, 10))
        self.assertEqual(events[0].dtend, datetime(2026, 9, 17, 11, 50))
        self.assertEqual(events[0].location, "教室 101")
        self.assertEqual(events[0].description, "")
        self.assertEqual(events[0].alarms, [15])

    def test_native_date_and_time_become_event(self) -> None:
        schedule = pd.DataFrame(
            {
                "课程编号": ["FT03P207"],
                "课程名称": ["金融风险管理"],
                "日期": [datetime(2026, 9, 17)],
                "开始时间": [time(10, 10)],
                "结束时间": [time(11, 50)],
                "地点": ["教室 101"],
                "备注": ["闭卷"],
            }
        )

        event = next(self.get_parser(schedule).get_exam(23000001))

        self.assertEqual(event.dtstart, datetime(2026, 9, 17, 10, 10))
        self.assertEqual(event.dtend, datetime(2026, 9, 17, 11, 50))
        self.assertEqual(event.description, "闭卷")

    def test_exams_use_requested_reminder(self) -> None:
        schedule = pd.DataFrame(
            {
                "课程编号": ["FT03P207"],
                "课程名称": ["金融风险管理"],
                "日期": [datetime(2026, 9, 17)],
                "开始时间": [time(10, 10)],
                "结束时间": [time(11, 50)],
            }
        )
        for minutes in (30, 0):
            with self.subTest(reminder_minutes=minutes):
                event = next(self.get_parser(schedule, minutes).get_exam(23000001))
                self.assertEqual(event.alarms, [minutes])

    def test_uploaded_workbook_mixed_dates_and_times(self) -> None:
        workbook = Workbook()
        students = workbook.active
        students.title = "补考学生名单"
        students.append(["学生名单"])
        students.append(["学号", "课程编号"])
        students.append([23000001, "FT03P207"])
        students.append([23000002, "OTHER"])
        schedule = workbook.create_sheet("总补辅导安排")
        schedule.append(["辅导安排"])
        schedule.append(["课程编号", "课程名称", "辅导日期", "辅导日期（备查）",
                         "辅导开始时间", "辅导结束时间", "辅导地点", "备注"])
        schedule.append(["FT03P207", "课程一", 46324, "10.29", time(15, 20),
                         "16；05", "教室101", "闭卷"])
        schedule.append(["FT03P207", "课程二", datetime(2026, 11, 4), "11.4",
                         "上午10点20", "上午11点", "教室102", None])
        schedule.append(["FT03P207", "无安排", "无辅导", "无", "无", "无"])
        schedule.append(["FT03P207", "空安排", None, None, None, None])
        schedule.append(["OTHER", "其他学生课程", 46324, "10.29", time(8), time(9)])
        content = BytesIO()
        workbook.save(content)
        workbook.close()

        events = list(ExamParser(content.getvalue(), 30).get_exam(23000001))

        self.assertEqual(len(events), 2)
        self.assertEqual(events[0].dtstart, datetime(2026, 10, 29, 15, 20))
        self.assertEqual(events[0].dtend, datetime(2026, 10, 29, 16, 5))
        self.assertEqual(events[0].location, "教室101")
        self.assertEqual(events[0].description, "闭卷")
        self.assertEqual(events[1].dtstart, datetime(2026, 11, 4, 10, 20))
        self.assertEqual(events[1].dtend, datetime(2026, 11, 4, 11))
        self.assertTrue(all(event.alarms == [30] for event in events))

    def test_split_times_work_with_iso_date(self) -> None:
        schedule = pd.DataFrame({
            "课程编号": ["FT03P207"], "课程名称": ["课程"],
            "日期": ["2026-10-29"], "开始时间": ["09:15"], "结束时间": ["10：00"],
        })
        event = next(self.get_parser(schedule).get_exam(23000001))
        self.assertEqual(event.dtstart, datetime(2026, 10, 29, 9, 15))
        self.assertEqual(event.dtend, datetime(2026, 10, 29, 10))

    def test_real_workbook_clock_formats(self) -> None:
        cases = [
            ("晚8：00", "晚9：30", time(20), time(21, 30)),
            ("12点", "13点", time(12), time(13)),
            (" 19：00", time(20, 30), time(19), time(20, 30)),
            ("英文辅导：10:00:00", "英文辅导：10:45:00", time(10), time(10, 45)),
            (14, 16, time(14), time(16)),
            (0.5, 0.5625, time(12), time(13, 30)),
        ]
        for start, end, expected_start, expected_end in cases:
            with self.subTest(start=start, end=end):
                schedule = pd.DataFrame({
                    "课程编号": ["FT03P207"], "日期": [46324],
                    "开始时间": [start], "结束时间": [end],
                })
                event = next(self.get_parser(schedule).get_exam(23000001))
                self.assertEqual(event.dtstart, datetime.combine(datetime(2026, 10, 29), expected_start))
                self.assertEqual(event.dtend, datetime.combine(datetime(2026, 10, 29), expected_end))

    def test_invalid_rows_do_not_create_current_time_events_or_block_valid_rows(self) -> None:
        schedule = pd.DataFrame({
            "课程编号": ["FT03P207"] * 8,
            "日期": [46324, 46324, 46324, "无辅导", None, "not a date", 46324, 46324],
            "开始时间": [None, "not a time", time(11), "无", None, time(9), "25:00", time(9)],
            "结束时间": [time(10), time(10), time(10), "无", None, time(10), time(10), time(10)],
        })
        with self.assertLogs("parser.exam", level="WARNING") as messages:
            events = list(self.get_parser(schedule).get_exam(23000001))
        self.assertEqual(len(messages.output), 5)
        self.assertEqual(len(events), 1)
        self.assertEqual(events[0].dtstart, datetime(2026, 10, 29, 9))
        self.assertEqual(events[0].dtend, datetime(2026, 10, 29, 10))

    def test_workbook_1904_epoch_is_used(self) -> None:
        workbook = Workbook()
        workbook.epoch = MAC_EPOCH
        students = workbook.active
        students.append(["学生名单"])
        students.append(["学号", "课程编号"])
        students.append([23000001, "FT03P207"])
        schedule = workbook.create_sheet("安排")
        schedule.append(["课程安排"])
        schedule.append(["课程编号", "日期", "开始时间", "结束时间"])
        schedule.append(["FT03P207", to_excel(datetime(2026, 10, 29), MAC_EPOCH), 14, 16])
        content = BytesIO()
        workbook.save(content)
        workbook.close()
        event = next(ExamParser(content.getvalue()).get_exam(23000001))
        self.assertEqual(event.dtstart, datetime(2026, 10, 29, 14))
        self.assertEqual(event.dtend, datetime(2026, 10, 29, 16))
